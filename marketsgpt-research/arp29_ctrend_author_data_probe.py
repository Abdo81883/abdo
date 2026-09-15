#!/usr/bin/env python3
import io, json, zipfile
from datetime import datetime, timezone
from pathlib import Path
import requests

ROOT=Path(__file__).resolve().parent
OUT=ROOT/"results";OUT.mkdir(parents=True,exist_ok=True)
SHEET_ID="1kE3P-4Q2KMRxLQxDbyRE2Gj2d2W-KDdG"
URLS=[
 ("xlsx",f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"),
 ("csv_gid0",f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=csv&gid=0"),
 ("gviz_csv",f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/gviz/tq?tqx=out:csv")
]
s=requests.Session();s.headers.update({"User-Agent":"MarketsGPT-CTREND-Probe/1.0"})
probes=[]
payload=None
for name,url in URLS:
    try:
        r=s.get(url,timeout=45,allow_redirects=True)
        item={"name":name,"url":url,"status":r.status_code,"contentType":r.headers.get("content-type"),"bytes":len(r.content),"finalUrl":r.url}
        if "text" in (r.headers.get("content-type") or "") or "csv" in (r.headers.get("content-type") or ""):
            item["head"]=r.text[:1000]
        probes.append(item)
        if r.status_code==200 and len(r.content)>100 and payload is None:
            payload=(name,r.content,r.headers.get("content-type") or "")
    except Exception as e:
        probes.append({"name":name,"url":url,"error":repr(e)})
audit={"schema":"mgpt_arp29_ctrend_author_data_probe_v1","generatedAtUtc":datetime.now(timezone.utc).isoformat(),"sheetId":SHEET_ID,"probes":probes,"accessible":payload is not None}
if payload:
    name,blob,ct=payload
    if name=="xlsx" or "spreadsheet" in ct:
        p=OUT/"ARP29_CTREND_AUTHOR_UPDATED_WEEKLY_FACTOR_DATA_20260915.xlsx";p.write_bytes(blob);audit["savedFile"]=p.name
        try:
            import openpyxl
            wb=openpyxl.load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
            audit["sheets"]=wb.sheetnames
            preview={}
            for ws in wb.worksheets:
                rows=[]
                for row in ws.iter_rows(min_row=1,max_row=12,values_only=True):
                    rows.append([None if v is None else str(v) for v in row[:12]])
                preview[ws.title]=rows
            audit["preview"]=preview
        except Exception as e:audit["xlsxInspectError"]=repr(e)
    else:
        p=OUT/"ARP29_CTREND_AUTHOR_UPDATED_WEEKLY_FACTOR_DATA_20260915.csv";p.write_bytes(blob);audit["savedFile"]=p.name;audit["previewText"]=blob[:3000].decode("utf-8","replace")
(OUT/"ARP29_CTREND_AUTHOR_DATA_PROBE_20260915.json").write_text(json.dumps(audit,indent=2,ensure_ascii=False)+"\n")
print(json.dumps(audit,indent=2,ensure_ascii=False))
